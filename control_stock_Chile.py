import flet as ft
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from openpyxl import Workbook
from bs4 import BeautifulSoup
import time
import threading
import os
import re

# Configuración
MAX_WORKERS = 20  # Número máximo de hilos
TIMEOUT = 10      # Tiempo de espera para solicitudes

headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
}

# Variables globales de control y para el Excel generado
processing_paused = False
processing_running = False
last_excel_file = None

def completar_codigo(codigo):
    """Agrega '001' a códigos de 8 dígitos."""
    if len(codigo) == 8:
        return codigo + "001"
    return codigo

def obtener_datos_producto(codigo_padre):
    """Extrae los datos del producto a partir de su código."""
    codigo_padre = completar_codigo(codigo_padre)
    url = f'https://www.marathon.cl/{codigo_padre}.html'
    status_web = 'N/A'
    try:
        start_time = time.time()
        response = requests.get(url, headers=headers, timeout=TIMEOUT, allow_redirects=True)
        load_time = time.time() - start_time
        final_url = response.url
        soup = BeautifulSoup(response.text, 'html.parser')
        
        # Determinar el estado web
        if 'marathon.cl/' not in final_url:
            status_web = 'REDIRECCIÓN DETECTADA'
        elif "PÁGINA NO ENCONTRADA" in soup.text.upper() or "EL PRODUCTO QUE BUSCAS NO EXISTE" in soup.text.upper():
            status_web = 'NO DISPONIBLE'
        elif response.status_code == 404:
            status_web = 'ERROR 404'
        else:
            status_web = 'DISPONIBLE'
        if final_url == 'https://www.marathon.cl/':
            status_web = 'REDIRECCIONADO AL INICIO'
        
        # PRECIO ACTUAL
        precio_actual = soup.select_one('span.value[content]')
        if not precio_actual:
            precio_actual = soup.select_one('span.price-sales')
        if precio_actual:
            if 'content' in precio_actual.attrs:
                precio_actual = precio_actual['content']
            else:
                precio_actual = precio_actual.text.strip()
        else:
            precio_actual = 'N/A'
        
        # FULL PRICE (columna C)
        full_price_elem = soup.select_one("#pdp del span span")
        if full_price_elem:
            full_price = full_price_elem.get_text(strip=True)
        else:
            full_price = 'N/A'
        
        # DESCUENTO (columna D)
        descuento_elem = soup.select_one('div.pd-item-promo')
        if descuento_elem:
            descuento_text = descuento_elem.text.strip()
            match = re.search(r'-?(\d+%?)', descuento_text)
            if match:
                descuento = match.group(1)
            else:
                descuento = descuento_text
        else:
            descuento = 'N/A'
        
        # CANTIDAD DE IMÁGENES
        imagenes = len(soup.select('img.galley_img'))
        
        # DESCRIPCIÓN COMERCIAL
        descripcion_element = soup.select_one('div.product-text[data-product-field="longDescription"]')
        if not descripcion_element:
            descripcion_element = soup.select_one('div.product-description')
        descripcion_texto = descripcion_element.get_text(strip=True) if descripcion_element else 'N/A'
        
        return (codigo_padre, precio_actual, full_price, descuento, imagenes, descripcion_texto, round(load_time, 2), url, status_web)
    
    except requests.exceptions.RequestException as e:
        print(f"ERROR DE RED AL PROCESAR {codigo_padre}: {str(e)}")
        return (codigo_padre, 'ERROR DE RED', 'N/A', 'N/A', 'N/A', 'N/A', 'N/A', url, 'ERROR DE RED')
    except Exception as e:
        print(f"ERROR INESPERADO AL PROCESAR {codigo_padre}: {str(e)}")
        return (codigo_padre, 'ERROR INESPERADO', 'N/A', 'N/A', 'N/A', 'N/A', 'N/A', url, 'ERROR INESPERADO')

def guardar_resultados(resultados):
    """Guarda los resultados en un archivo Excel y retorna su nombre."""
    global last_excel_file
    wb = Workbook()
    ws = wb.active
    ws.title = "DATOS PRODUCTOS"
    headers_excel = [
        "CODIGO", "PRECIO ACTUAL", "FULL PRICE", "DESCUENTO",
        "CANT. IMG", "DESCRIPCION COMERCIAL", "TIEMPO DE CARGA (S)", "URL", "CONTROL STOCK"
    ]
    for col, header in enumerate(headers_excel, start=1):
        ws.cell(row=1, column=col, value=header)
    for i, row in enumerate(resultados, start=2):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)
    fecha_hora_actual = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    nombre_archivo = f"DATOS_PRODUCTOS_WEB_{fecha_hora_actual}.xlsx"
    wb.save(nombre_archivo)
    last_excel_file = nombre_archivo
    return nombre_archivo

def process_codes(codes, progress_callback, status_callback):
    """Procesa la lista de códigos y actualiza el progreso y estado."""
    global processing_paused
    results = []
    total = len(codes)
    start_time = datetime.now()
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {executor.submit(obtener_datos_producto, code): code for code in codes}
        for i, future in enumerate(as_completed(futures), start=1):
            if processing_paused:
                status_callback("PROCESO PAUSADO.")
                # Se intenta cancelar los futuros pendientes
                for fut in futures:
                    fut.cancel()
                break
            result = future.result()
            results.append(result)
            elapsed = datetime.now() - start_time
            elapsed_str = str(elapsed).split('.')[0]
            status_callback(f"PROCESANDO CÓDIGO {i}/{total} - TIEMPO TRANSCURRIDO: {elapsed_str}")
            progress_callback(i / total)
            time.sleep(0.5)
    return results

def run_processing(codes, progress_callback, status_callback, done_callback):
    """Ejecuta el procesamiento en segundo plano, guardando resultados parciales si se pausa."""
    global processing_running, processing_paused
    processing_running = True
    processing_paused = False
    results = process_codes(codes, progress_callback, status_callback)
    if results:
        file_name = guardar_resultados(results)
        if processing_paused:
            status_callback(f"PROCESAMIENTO PAUSADO. RESULTADOS PARCIALES GUARDADOS EN: {file_name}")
        else:
            status_callback(f"PROCESAMIENTO COMPLETADO. ARCHIVO GUARDADO: {file_name}")
    processing_running = False
    done_callback()

def main(page: ft.Page):
    page.bgcolor = ft.colors.WHITE
    page.title = "SCRAPER DE MARATHON.CL"
    page.vertical_alignment = ft.MainAxisAlignment.START
    page.horizontal_alignment = ft.CrossAxisAlignment.CENTER

    # Campo de códigos responsivo
    txt_codes = ft.TextField(
        label="CÓDIGOS (SEPARADOS POR ESPACIO O SALTO DE LÍNEA)",
        multiline=True,
        expand=True,
        border_color=ft.colors.GREY,
        border_width=2,
        content_padding=10
    )
    
    file_status = ft.Text(value="Ningún archivo cargado", color=ft.colors.BLACK)
    status_text = ft.Text(value="ESTADO: ESPERANDO INICIAR...", color=ft.colors.BLACK)
    
    progress_bar = ft.ProgressBar(
        value=0,
        color=ft.colors.BLUE,
        bgcolor=ft.colors.LIGHT_BLUE,
        height=20,
        width=600
    )
    
    # Botones con estilo dinámico
    btn_load_file = ft.ElevatedButton(
        "CARGAR ARCHIVO TXT",
        bgcolor=ft.colors.BLUE,
        color=ft.colors.WHITE,
        style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=8))
    )
    btn_start = ft.ElevatedButton(
        "INICIAR PROCESAMIENTO",
        bgcolor=ft.colors.BLUE,
        color=ft.colors.WHITE,
        style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=8))
    )
    btn_pause = ft.ElevatedButton(
        "PAUSAR PROCESAMIENTO",
        bgcolor=ft.colors.BLUE,
        color=ft.colors.WHITE,
        style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=8))
    )
    btn_open_excel = ft.ElevatedButton(
        "ABRIR EXCEL",
        bgcolor=ft.colors.BLUE,
        color=ft.colors.WHITE,
        style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=8))
    )
    btn_pause.disabled = True

    # Configuración del FilePicker (se leerá por ruta)
    file_picker = ft.FilePicker(on_result=lambda e: file_picker_result(e, txt_codes, file_status, page))
    page.overlay.append(file_picker)
    
    processing_thread = None
    
    def update_progress(value):
        progress_bar.value = value
        page.update()
        
    def update_status(msg):
        status_text.value = msg
        page.update()
        
    def processing_done():
        btn_start.disabled = False
        btn_pause.disabled = True
        progress_bar.value = 0
        page.update()
    
    def on_start_click(e):
        nonlocal processing_thread
        if txt_codes.value.strip() == "":
            update_status("POR FAVOR, INGRESA O CARGA CÓDIGOS.")
            return
        codes = [code.strip() for code in txt_codes.value.replace("\n", " ").split() if code.strip()]
        btn_start.disabled = True
        btn_pause.disabled = False
        page.update()
        processing_thread = threading.Thread(
            target=run_processing,
            args=(codes, update_progress, update_status, processing_done),
            daemon=True
        )
        processing_thread.start()
    
    def on_pause_click(e):
        global processing_paused
        processing_paused = True
        update_status("PAUSANDO PROCESO, SE GUARDARÁN LOS RESULTADOS PARCIALES...")
        btn_pause.disabled = True
        page.update()
    
    def on_open_excel_click(e):
        if last_excel_file and os.path.exists(last_excel_file):
            try:
                os.startfile(last_excel_file)
            except Exception as ex:
                update_status(f"ERROR AL ABRIR EL ARCHIVO: {ex}")
        else:
            update_status("NO SE HA GENERADO NINGÚN ARCHIVO EXCEL.")
    
    btn_start.on_click = on_start_click
    btn_pause.on_click = on_pause_click
    btn_load_file.on_click = lambda e: file_picker.pick_files(allow_multiple=False)
    btn_open_excel.on_click = on_open_excel_click
    
    page.add(
        ft.Container(content=txt_codes, expand=True, padding=10),
        file_status,
        ft.Row(
            controls=[btn_load_file, btn_start, btn_pause, btn_open_excel],
            alignment=ft.MainAxisAlignment.CENTER,
            spacing=10,
            run_spacing=10,
            wrap=True
        ),
        ft.Container(content=progress_bar, padding=10),
        status_text
    )

def file_picker_result(e, txt_codes, file_status, page):
    if e.files:
        file = e.files[0]
        try:
            with open(file.path, "r", encoding="utf-8") as f:
                content = f.read()
            txt_codes.value = content
            file_status.value = f"ARCHIVO CARGADO: {file.name}"
            page.update()
        except Exception as ex:
            file_status.value = "ERROR AL LEER EL ARCHIVO"
            page.update()
            print("ERROR AL LEER EL ARCHIVO:", ex)

ft.app(target=main)
