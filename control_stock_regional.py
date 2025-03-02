import flet as ft
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from openpyxl import Workbook
from bs4 import BeautifulSoup
import time
import threading
import os
import re
import sys
import subprocess
import asyncio
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout

# Variables globales
processing_paused = False
processing_running = False
last_excel_file = None

results_data = []  # Acumula los registros extraídos
start_time = None
total_codes = 0
processed_codes = 0

# -------------------------------------------------------
# Bitácora (log) en la interfaz
log_area = None  # Se asignará en main()

def log_message(msg: str):
    """Agrega un mensaje al log_area, con salto de línea."""
    if log_area:
        log_area.value += f"{msg}\n"
        log_area.update()
    print(msg)  # También en consola

# -------------------------------------------------------
async def extraer_datos_playwright(codigo, pais):
    """
    Extrae datos dinámicos de un producto usando Playwright y BeautifulSoup.
    - Determina si hay alguna talla habilitada (stock_status).
    - Extrae precio (columna C), descuento (columna H) y full price (columna I)
      con los selectores que especificaste.
    - Si ocurre un timeout en page.goto, se captura la excepción y se retorna
      un estado de "TIMEOUT".
    """
    if len(codigo) == 8:
        codigo += "001"
    base_url = f"https://www.marathon.store/{pais}/p/{codigo}"
    inicio = time.time()

    try:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            page = await browser.new_page()

            # Aumentamos el timeout a 60s
            page.set_default_navigation_timeout(60000)

            try:
                await page.goto(base_url, wait_until="load")
            except PlaywrightTimeout:
                log_message(f"[TIMEOUT] {codigo} excedió el tiempo de carga.")
                return (codigo, "TIMEOUT", "N/A", 0, "N/A", base_url, round(time.time()-inicio,2), "N/A", "N/A")

            # Intentamos esperar la tabla de tallas, pero si no aparece, no bloqueamos
            try:
                await page.wait_for_selector("table.table-size-selector", timeout=10000)
            except PlaywrightTimeout:
                log_message(f"[NO TABLE] {codigo} no cargó la tabla de tallas en 10s.")

            content = await page.content()
            await browser.close()

    except Exception as e:
        log_message(f"[ERROR] No se pudo cargar {base_url}: {e}")
        return (codigo, "ERROR AL CARGAR", "N/A", 0, "N/A", base_url, 0, "N/A", "N/A")

    fin = time.time()
    tiempo_respuesta = round(fin - inicio, 2)
    soup = BeautifulSoup(content, "html.parser")
    
    # --- CONTROL DE STOCK: ver si existe alguna talla disponible ---
    table = soup.find("table", class_=lambda x: x and "table-size-selector" in x)
    if table:
        cells = table.find_all("td")
        available = False
        valid_size_found = False
        for cell in cells:
            span = cell.find("span")
            if span:
                valid_size_found = True
                classes = cell.get("class", [])
                # Si el <td> NO contiene "pdp-size-disabled", se considera disponible
                if "pdp-size-disabled" not in classes:
                    available = True
                    break
        if not valid_size_found:
            stock_status = "SIN TALLAS"
        else:
            stock_status = "DISPONIBLE" if available else "AGOTADO"
    else:
        stock_status = "TABLA DE TALLAS NO ENCONTRADA"
    
    # --- PRECIO (Columna C) ---
    try:
        precio_elem = soup.select_one(
            "main > div:nth-of-type(4) > div:nth-of-type(1) > div:nth-of-type(1) > "
            "div:nth-of-type(4) > div > div:nth-of-type(1) > div:nth-of-type(2) > "
            "div:nth-of-type(2) > div:nth-of-type(1)"
        )
        precio = precio_elem.get_text(strip=True) if precio_elem else "N/A"
    except Exception as e:
        log_message(f"[PRECIO FAIL] {codigo}: {e}")
        precio = "N/A"
    
    # --- FULL PRICE (Columna I) ---
    try:
        full_price_elem = soup.select_one(
            "main > div:nth-of-type(4) > div:nth-of-type(1) > div:nth-of-type(1) > "
            "div:nth-of-type(1) > div:nth-of-type(3) > div:nth-of-type(2) > "
            "div:nth-of-type(1) > del"
        )
        full_price = full_price_elem.get_text(strip=True) if full_price_elem else "N/A"
    except Exception as e:
        log_message(f"[FULL PRICE FAIL] {codigo}: {e}")
        full_price = "N/A"
    
    # --- DESCUENTO (Columna H) ---
    descuento = "N/A"
    try:
        descuento_elem = soup.select_one(
            "main > div:nth-of-type(4) > div:nth-of-type(1) > div:nth-of-type(1) > "
            "div:nth-of-type(1) > div:nth-of-type(3) > div:nth-of-type(1) > p"
        )
        if descuento_elem:
            desc_text = descuento_elem.get_text(strip=True)
            match = re.search(r"(\d+%)", desc_text)
            if match:
                descuento = match.group(1)
            else:
                descuento = desc_text
    except Exception as e:
        log_message(f"[DESCUENTO FAIL] {codigo}: {e}")
    
    # --- IMÁGENES ---
    imagenes = soup.select("div.desktop-image-gallery img")
    enlaces_img = []
    for img in imagenes:
        src = img.get("src") or img.get("data-src")
        if src:
            enlaces_img.append(src)
    cant_img = len(enlaces_img)
    enlaces_str = ", ".join(enlaces_img)
    
    return (
        codigo,         # A
        stock_status,   # B
        precio,         # C
        cant_img,       # D
        enlaces_str,    # E
        base_url,       # F
        tiempo_respuesta,# G
        descuento,      # H
        full_price      # I
    )

def extraer_datos_playwright_sync(codigo, pais):
    return asyncio.run(extraer_datos_playwright(codigo, pais))

def process_codes(codes, pais, progress_callback, status_callback, done_callback):
    global processing_paused, processing_running, results_data
    global total_codes, processed_codes, start_time

    processing_running = True
    processing_paused = False
    results_data = []
    total_codes = len(codes)
    processed_codes = 0
    start_time = datetime.now()
    
    max_workers = 5
    futures = {}

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {executor.submit(extraer_datos_playwright_sync, code, pais): code for code in codes}
        try:
            for future in as_completed(futures):
                if processing_paused:
                    status_callback("PROCESO PAUSADO. CANCELANDO tareas pendientes y guardando resultados...")
                    for f in futures:
                        if not f.done():
                            f.cancel()
                    break

                result = future.result()
                results_data.append(result)
                processed_codes += 1

                progress_callback(processed_codes / total_codes)
                elapsed = datetime.now() - start_time
                h, r = divmod(elapsed.seconds, 3600)
                m, s = divmod(r, 60)
                status_callback(
                    f"Procesando {processed_codes}/{total_codes} - "
                    f"Tiempo {h:02d}:{m:02d}:{s:02d}"
                )
                log_message(f"[OK] {result[0]} -> STOCK={result[1]} | PRECIO={result[2]}")
                time.sleep(0.2)

        except Exception as e:
            status_callback(f"Error durante el procesamiento: {e}")
    
    file_name = guardar_resultados()
    if processing_paused:
        status_callback(f"RESULTADOS PARCIALES GUARDADOS: {file_name}")
    else:
        status_callback(f"PROCESAMIENTO COMPLETADO. ARCHIVO: {file_name}")

    processing_running = False
    done_callback()

def guardar_resultados():
    global results_data, last_excel_file
    wb = Workbook()
    ws = wb.active
    ws.title = "RESULTADOS"

    headers = [
        "CÓDIGO", "CONTROL STOCK", "PRECIO", "CANT. IMG", "ENLACES IMÁGENES",
        "URL", "TIEMPO (s)", "DESCUENTO", "FULL PRICE"
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)

    for i, row in enumerate(results_data, start=2):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)

    fecha_hora = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre_archivo = f"RESULTADOS_{fecha_hora}.xlsx"
    wb.save(nombre_archivo)
    last_excel_file = nombre_archivo
    log_message(f"[EXCEL] Guardado: {nombre_archivo}")
    return nombre_archivo

def abrir_excel():
    global last_excel_file
    if last_excel_file and os.path.exists(last_excel_file):
        try:
            if sys.platform.startswith("win"):
                os.startfile(last_excel_file)
            elif sys.platform.startswith("darwin"):
                subprocess.Popen(["open", last_excel_file])
            else:
                subprocess.Popen(["xdg-open", last_excel_file])
        except Exception as ex:
            return f"Error al abrir el archivo: {ex}"
    else:
        return "No se ha generado ningún archivo Excel todavía."
    return "Archivo abierto."

def main(page: ft.Page):
    page.title = "Scraper Marathon Store (Playwright + BS4)"
    page.bgcolor = ft.colors.WHITE
    page.vertical_alignment = ft.MainAxisAlignment.START
    page.scroll = ft.ScrollMode.AUTO

    global log_area
    # Este Text contendrá los mensajes del log
    log_area = ft.Text(value="", selectable=True, color=ft.colors.BLACK)

    # Lo ponemos dentro de una Column con scroll
    log_column = ft.Column(
        controls=[log_area],
        width=600,
        height=200,
        scroll=ft.ScrollMode.AUTO
    )

    # Y si quieres un borde, lo envuelves en un Container
    log_container = ft.Container(
        content=log_column,
        border=ft.border.all(1, ft.colors.GREY),
        border_radius=5,
        padding=5
    )

    status_text = ft.Text(value="ESTADO: Esperando inicio...", color=ft.colors.BLACK)
    progress_bar = ft.ProgressBar(width=600, value=0)

    txt_codes = ft.TextField(
        label="Códigos (separados por espacio o salto de línea)",
        multiline=True,
        width=600,
        height=150,
        border_color=ft.colors.GREY,
        border_width=2
    )
    dd_pais = ft.Dropdown(
        label="País",
        options=[
            ft.dropdown.Option(key="pe", text="Perú"),
            ft.dropdown.Option(key="bo", text="Bolivia"),
            ft.dropdown.Option(key="ec", text="Ecuador")
        ],
        width=200
    )
    file_picker = ft.FilePicker(on_result=lambda e: on_file_picked(e, txt_codes))
    page.overlay.append(file_picker)

    btn_load_file = ft.ElevatedButton(
        text="Cargar archivo TXT",
        on_click=lambda _: file_picker.pick_files(allow_multiple=False),
        bgcolor=ft.colors.BLUE, color=ft.colors.WHITE
    )
    btn_start = ft.ElevatedButton(
        text="Iniciar",
        bgcolor=ft.colors.GREEN, color=ft.colors.WHITE
    )
    btn_pause = ft.ElevatedButton(
        text="Pausar",
        bgcolor=ft.colors.RED, color=ft.colors.WHITE
    )
    btn_open_excel = ft.ElevatedButton(
        text="Abrir Excel",
        bgcolor=ft.colors.BLUE, color=ft.colors.WHITE
    )

    btn_pause.disabled = True
    btn_open_excel.disabled = True

    def update_progress(value: float):
        progress_bar.value = value
        page.update()

    def update_status(msg: str):
        status_text.value = msg
        page.update()

    def processing_done():
        btn_start.disabled = False
        btn_pause.disabled = True
        progress_bar.value = 0
        # Habilitamos abrir Excel
        btn_open_excel.disabled = False
        page.update()

    def on_start_click(e):
        if not dd_pais.value:
            update_status("Por favor, selecciona un país.")
            return
        codes_str = txt_codes.value.strip()
        if not codes_str:
            update_status("Por favor, ingresa o carga algunos códigos.")
            return
        codes = [c.strip() for c in codes_str.replace("\n", " ").split() if c.strip()]
        btn_start.disabled = True
        btn_pause.disabled = False
        btn_open_excel.disabled = True
        update_status("Iniciando procesamiento...")
        page.update()

        # Iniciamos en un hilo para no bloquear la UI
        threading.Thread(
            target=process_codes,
            args=(codes, dd_pais.value, update_progress, update_status, processing_done),
            daemon=True
        ).start()

    def on_pause_click(e):
        global processing_paused
        processing_paused = True
        update_status("Pausando el proceso y guardando resultados...")
        btn_pause.disabled = True
        # Permite abrir Excel parcial
        btn_open_excel.disabled = False
        page.update()

    def on_open_excel_click(e):
        msg = abrir_excel()
        update_status(msg)

    btn_start.on_click = on_start_click
    btn_pause.on_click = on_pause_click
    btn_open_excel.on_click = on_open_excel_click

    page.add(
        ft.Column([
            ft.Row([dd_pais]),
            txt_codes,
            ft.Row([btn_load_file, btn_start, btn_pause, btn_open_excel]),
            progress_bar,
            status_text,
            ft.Text("Bitácora (Log):", color=ft.colors.BLACK, weight=ft.FontWeight.BOLD),
            log_container
        ], spacing=10)
    )

def on_file_picked(e: ft.FilePickerResultEvent, txt_codes: ft.TextField):
    if e.files:
        file = e.files[0]
        try:
            with open(file.path, "r", encoding="utf-8") as f:
                contenido = f.read()
            txt_codes.value = contenido
            txt_codes.update()
            log_message(f"Archivo {file.name} cargado.")
        except Exception as ex:
            log_message(f"Error al leer el archivo: {ex}")

# Ejecutar la app
ft.app(target=main)
